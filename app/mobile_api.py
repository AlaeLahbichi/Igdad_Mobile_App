"""API de l'application mobile Igdad (réseau local).

    bash scripts/lake_start.sh                  # MinIO : catalogue, photos, chants, tests enregistrés
    .venv/bin/python app/mobile_api.py          # écoute sur 0.0.0.0:8000 (le téléphone doit être sur le même Wi-Fi)

Catalogue      GET  /api/m/health | /api/m/species | /api/m/species/<espece>
Médias         GET  /api/m/photo/<espece>?w=  /api/m/gallery/<espece>/<i>?w=  /api/m/song/<espece>/<chant|cri>
Audio          POST /api/m/audio/stream   {session, sr, pcm}         pcm = int16 LE en base64 (nouveaux échantillons)
               POST /api/m/audio/file     {session, name, data}      fichier audio en base64 (wav, mp3, m4a, ...)
Image          POST /api/m/image          {session, mode, data}      JPEG en base64
               (les trois acceptent aussi lat/lon : position du téléphone au moment de l'analyse)
Tests          POST /api/m/save {session} | GET /api/m/tests | GET /api/m/tests/<id>/media | /api/m/tests/<id>
Dashboard      GET  /api/m/dashboard
Réglages       GET/POST /api/m/settings  {threshold}   seuil modifiable depuis l'app (data/igdad/settings.json)

Seuil de confiance : 90 % par défaut, réglable depuis l'app (50 à 99 %). Chaque espèce détectée au-dessus du seuil est ajoutée UNE fois par session
(écoute ou analyse) à data/igdad/detections.csv et data/igdad/detections.xlsx (feuilles Détections + Compteur),
et son compteur augmente de 1 dans data/igdad/compteur_especes.xlsx / .csv (les 83 espèces, count).
Les tests enregistrés vont dans taghrid-bronze/igdad_tests/<audio|image>/.
"""
import argparse
import base64
import csv
import io
import json
import re
import subprocess
import sys
import tempfile
import threading
import time
import uuid
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

import librosa
import numpy as np
import soundfile as sf
import torch
from PIL import Image, ImageOps

sys.path.insert(0, str(Path(__file__).resolve().parent))
from server import ROOT, Predictor, s3_client  # noqa: E402

BRONZE, GOLD = "taghrid-bronze", "taghrid-gold"
THRESHOLD = 0.90                 # valeur par défaut ; la valeur courante est dans data/igdad/settings.json
THRESHOLD_RANGE = (0.50, 0.99)
# Temps réel : garde-fous contre les fausses détections
MIN_AUDIO_S = 3.0                # pas de prédiction avant 3 s d'écoute (le modèle n'a vu que des segments >= 2 s)
MIN_LEVEL = 0.003                # RMS (sur 0,25 s) du passage le plus fort de la fenêtre, ~ -50 dBFS
CONFIRM_WINDOWS = 2              # même espèce en tête, au-dessus du seuil, sur N prédictions consécutives
TESTS_PREFIX = "igdad_tests/"
EXPORT_DIR = ROOT / "data" / "igdad"
SESSION_TTL = 3600
SETTINGS_FILE = EXPORT_DIR / "settings.json"


def load_settings():
    global THRESHOLD
    try:
        THRESHOLD = float(json.loads(SETTINGS_FILE.read_text())["threshold"])
    except (FileNotFoundError, KeyError, ValueError):
        pass


def set_threshold(value):
    global THRESHOLD
    try:
        value = round(float(value), 2)
    except (TypeError, ValueError):
        raise ValueError("seuil invalide") from None
    if not THRESHOLD_RANGE[0] <= value <= THRESHOLD_RANGE[1]:
        raise ValueError(f"le seuil doit être entre {THRESHOLD_RANGE[0]:.0%} et {THRESHOLD_RANGE[1]:.0%}")
    THRESHOLD = value
    SETTINGS_FILE.parent.mkdir(parents=True, exist_ok=True)
    SETTINGS_FILE.write_text(json.dumps({"threshold": THRESHOLD}))
    return THRESHOLD
MAX_SESSION_S = 15 * 60


# ---------------------------------------------------------------- modèles ----------------------------------
class AudioModel(Predictor):
    def windows(self, y, sr):
        """Fenêtres de 5 s (pas de 2,5 s) sur tout le fichier -> probabilités (n, 83) + débuts en secondes."""
        c = self.cfg
        y = librosa.resample(y.astype(np.float32), orig_sr=sr, target_sr=c["sr"]) if sr != c["sr"] else y
        n, hop = int(c["chunk_s"] * c["sr"]), int(c["chunk_s"] * c["sr"] / 2)
        if len(y) < n:
            y = np.pad(y, (0, n - len(y)))
        starts = list(range(0, len(y) - n + 1, hop))
        if (len(y) - n) % hop:
            starts.append(len(y) - n)
        feats = []
        for s in starts:
            m = librosa.feature.melspectrogram(y=y[s:s + n], sr=c["sr"], n_fft=c["n_fft"], hop_length=c["hop_length"],
                                               n_mels=c["n_mels"], fmin=c["fmin"], fmax=c["fmax"], power=2.0)
            db = librosa.power_to_db(m, ref=np.max, top_db=80.0).astype(np.float16).astype(np.float32)
            feats.append((db + 40.0) / 40.0)
        x = torch.from_numpy(np.stack(feats))[:, None]
        probs = []
        with self.lock, torch.no_grad():
            for i in range(0, len(x), 32):
                probs.append(self.model(x[i:i + 32]).softmax(-1).numpy())
        return np.concatenate(probs), [s / c["sr"] for s in starts]


class ImageModel:
    def __init__(self, weights):
        from ultralytics import YOLO
        self.model = YOLO(str(weights))
        self.names = self.model.names
        self.device = "mps" if torch.backends.mps.is_available() else "cpu"
        self.lock = threading.Lock()

    def detect(self, data):
        im = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
        with self.lock:
            res = self.model.predict(im, conf=0.25, iou=0.5, imgsz=640, device=self.device, verbose=False)[0]
        boxes = [{"species": self.names[int(c)], "conf": float(p), "box": [round(v, 4) for v in b.tolist()]}
                 for b, c, p in zip(res.boxes.xyxyn, res.boxes.cls, res.boxes.conf)]
        return im, sorted(boxes, key=lambda b: -b["conf"])


# ---------------------------------------------------------------- journal CSV + Excel ----------------------
class DetectionLog:
    COLS = ["date", "heure", "mode", "source", "session", "nom_francais", "nom_scientifique", "famille",
            "confiance_pct", "debut_s", "fin_s", "latitude", "longitude"]

    def __init__(self, catalog):
        from openpyxl import Workbook, load_workbook
        self.catalog, self.lock = catalog, threading.Lock()
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        self.csv, self.xlsx = EXPORT_DIR / "detections.csv", EXPORT_DIR / "detections.xlsx"
        if not self.csv.exists():
            with self.csv.open("w", newline="", encoding="utf-8-sig") as f:
                csv.writer(f, delimiter=";").writerow(self.COLS)
        else:
            self._migrate_csv()
        if self.xlsx.exists():
            self.wb = load_workbook(self.xlsx)
            ws = self.wb["Détections"]
            for i, col in enumerate(self.COLS, 1):   # anciennes versions : colonnes GPS absentes
                if ws.cell(1, i).value != col:
                    ws.cell(1, i).value = col
            self._style(ws)
        else:
            self.wb = Workbook()
            ws = self.wb.active
            ws.title = "Détections"
            ws.append(self.COLS)
            self.wb.create_sheet("Compteur")
            self._style(ws)
            self._counter()
            self.wb.save(self.xlsx)

    def _migrate_csv(self):
        with self.csv.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f, delimiter=";"))
        if rows and rows[0] != self.COLS:
            with self.csv.open("w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f, delimiter=";")
                w.writerow(self.COLS)
                for r in rows[1:]:
                    w.writerow((r + [""] * len(self.COLS))[:len(self.COLS)])

    def rows(self):
        """Journal complet (dictionnaires), lu depuis le CSV."""
        with self.lock, self.csv.open(encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f, delimiter=";"))

    def _style(self, ws):
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F5E4B")
        for col, w in zip("ABCDEFGHIJKLM", (11, 9, 8, 11, 10, 28, 28, 18, 13, 9, 9, 11, 11)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

    def _counter(self):
        from openpyxl.styles import Font, PatternFill
        det, ws = self.wb["Détections"], self.wb["Compteur"]
        counts = {}
        for r in det.iter_rows(min_row=2, values_only=True):
            if r[6]:
                c = counts.setdefault(r[6], {"fr": r[5], "audio": 0, "image": 0, "last": ""})
                c["audio" if r[2] == "Audio" else "image"] += 1
                c["last"] = max(c["last"], f"{r[0]} {r[1]}")
        ws.delete_rows(1, ws.max_row)
        ws.append(["nom_francais", "nom_scientifique", "detections_audio", "detections_image", "total",
                   "derniere_detection"])
        for sci, c in sorted(counts.items(), key=lambda kv: -(kv[1]["audio"] + kv[1]["image"])):
            ws.append([c["fr"], sci, c["audio"], c["image"], c["audio"] + c["image"], c["last"]])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F5E4B")
        for col, w in zip("ABCDEF", (28, 28, 17, 17, 8, 20)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

    def add(self, mode, source, session, species, conf, start=None, end=None, pos=None):
        e = self.catalog.get(species, {})
        now = datetime.now()
        lat, lon = pos or (None, None)
        row = [now.strftime("%Y-%m-%d"), now.strftime("%H:%M:%S"), mode, source, session[:8], e.get("fr") or "",
               e.get("sci", species.replace("_", " ")), e.get("family_fr") or e.get("family") or "",
               round(conf * 100, 1), None if start is None else round(start, 1), None if end is None else round(end, 1),
               lat, lon]
        with self.lock:
            with self.csv.open("a", newline="", encoding="utf-8-sig") as f:
                csv.writer(f, delimiter=";").writerow(["" if v is None else v for v in row])
            self.wb["Détections"].append(row)
            self._counter()
            try:
                self.wb.save(self.xlsx)
            except PermissionError:
                print("detections.xlsx est verrouillé (ouvert ailleurs ?) : la ligne est dans le CSV.")
        self.counter.increment(species, mode, f"{row[0]} {row[1]}")


class SpeciesCounter:
    """compteur_especes.xlsx / .csv : les 83 espèces avec leur nombre de détections (+1 à chaque détection)."""
    COLS = ["nom_francais", "nom_scientifique", "famille", "count", "count_audio", "count_image", "derniere_detection"]

    def __init__(self, catalog, journal_rows):
        self.lock = threading.Lock()
        self.csv, self.xlsx = EXPORT_DIR / "compteur_especes.csv", EXPORT_DIR / "compteur_especes.xlsx"
        self.data = {sp: {"fr": e.get("fr") or "", "sci": e["sci"], "family": e.get("family_fr") or e.get("family") or "",
                          "audio": 0, "image": 0, "last": ""}
                     for sp, e in sorted(catalog.items(), key=lambda kv: (kv[1].get("fr") or kv[0]).lower())}
        by_sci = {v["sci"]: sp for sp, v in self.data.items()}
        if self.csv.exists():       # reprise du compteur existant
            with self.csv.open(encoding="utf-8-sig", newline="") as f:
                for r in csv.DictReader(f, delimiter=";"):
                    sp = by_sci.get(r["nom_scientifique"])
                    if sp:
                        self.data[sp].update(audio=int(r["count_audio"] or 0), image=int(r["count_image"] or 0),
                                             last=r["derniere_detection"])
        else:                       # première fois : on part du journal existant
            for r in journal_rows:
                sp = by_sci.get(r["nom_scientifique"])
                if sp:
                    d = self.data[sp]
                    d["audio" if r["mode"] == "Audio" else "image"] += 1
                    d["last"] = max(d["last"], f"{r['date']} {r['heure']}")
        self._write()

    def increment(self, species, mode, when):
        with self.lock:
            d = self.data.get(species)
            if d is None:
                return
            d["audio" if mode == "Audio" else "image"] += 1
            d["last"] = when
            self._write()

    def table(self):
        return [{"species": sp, **d, "count": d["audio"] + d["image"]} for sp, d in self.data.items()]

    def _write(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        rows = [[d["fr"], d["sci"], d["family"], d["audio"] + d["image"], d["audio"], d["image"], d["last"]]
                for d in self.data.values()]
        with self.csv.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(self.COLS)
            w.writerows(rows)
        wb = Workbook()
        ws = wb.active
        ws.title = "Compteur espèces"
        ws.append(self.COLS)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F5E4B")
        for col, w in zip("ABCDEFG", (28, 28, 22, 8, 12, 12, 20)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:G{len(rows) + 1}"
        try:
            wb.save(self.xlsx)
        except PermissionError:
            print("compteur_especes.xlsx est verrouillé (ouvert ailleurs ?) : le CSV est à jour.")


# ---------------------------------------------------------------- sessions ---------------------------------
class Sessions:
    def __init__(self):
        self.data, self.lock = {}, threading.Lock()

    def get(self, sid, kind, source):
        if not re.fullmatch(r"[A-Za-z0-9-]{8,64}", sid or ""):
            raise ValueError("session invalide")
        with self.lock:
            now = time.time()
            for k in [k for k, v in self.data.items() if now - v["touched"] > SESSION_TTL]:
                del self.data[k]
            s = self.data.setdefault(sid, {"kind": kind, "source": source, "created": datetime.now().isoformat(
                timespec="seconds"), "audio": [], "sr": None, "samples": 0, "image": None, "boxes": [],
                "detections": {}, "name": None, "streak": (None, 0)})
            s["touched"] = now
            return s


def merge_detection(s, species, conf, start=None, end=None):
    """Met à jour la détection agrégée de la session ; renvoie True si l'espèce est nouvelle."""
    d = s["detections"].get(species)
    if d is None:
        s["detections"][species] = {"species": species, "conf": conf, "count": 1, "segments": []}
        d, new = s["detections"][species], True
    else:
        d["conf"], d["count"], new = max(d["conf"], conf), d["count"] + 1, False
    if start is not None:
        if d["segments"] and start <= d["segments"][-1][1] + 0.01:
            d["segments"][-1][1] = max(d["segments"][-1][1], end)
        else:
            d["segments"].append([round(start, 2), round(end, 2)])
    return new


def position(req):
    """lat/lon envoyés par le téléphone (facultatifs) -> (lat, lon) arrondis ou None."""
    try:
        lat, lon = float(req["lat"]), float(req["lon"])
    except (KeyError, TypeError, ValueError):
        return None
    if -90 <= lat <= 90 and -180 <= lon <= 180:
        return round(lat, 5), round(lon, 5)
    return None


def detections_list(s):
    return sorted(s["detections"].values(), key=lambda d: -d["conf"])


def decode_audio(data, name):
    """Décode wav/flac/mp3/ogg via libsndfile, sinon m4a/aac/3gp/caf via afconvert (macOS)."""
    try:
        y, sr = sf.read(io.BytesIO(data), dtype="float32", always_2d=True)
        return y.mean(1), sr
    except Exception:
        pass
    ext = Path(name or "audio.m4a").suffix.lower() or ".m4a"
    if not re.fullmatch(r"\.[a-z0-9]{1,5}", ext):
        ext = ".m4a"
    with tempfile.TemporaryDirectory() as d:
        src, dst = Path(d) / f"in{ext}", Path(d) / "out.wav"
        src.write_bytes(data)
        subprocess.run(["afconvert", "-f", "WAVE", "-d", "LEI16", str(src), str(dst)], check=True,
                       capture_output=True, timeout=120)
        y, sr = sf.read(dst, dtype="float32", always_2d=True)
        return y.mean(1), sr


# ---------------------------------------------------------------- HTTP --------------------------------------
def make_handler(ctx):
    s3, catalog, audio_model, image_model, sessions, log = (ctx[k] for k in (
        "s3", "catalog", "audio", "image", "sessions", "log"))
    thumbs, thumbs_lock = {}, threading.Lock()

    def bronze(key):
        return s3.get_object(Bucket=BRONZE, Key=key)["Body"].read()

    def thumb(key, w):
        w = max(64, min(int(w or 0), 1600)) if w else 0
        ck = (key, w)
        with thumbs_lock:
            if ck in thumbs:
                return thumbs[ck]
        data = bronze(key)
        if w:
            im = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
            im.thumbnail((w, w))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=82)
            data = buf.getvalue()
        with thumbs_lock:
            if len(thumbs) > 600:
                thumbs.clear()
            thumbs[ck] = data
        return data

    summary = [{k: e[k] for k in ("species", "sci", "fr", "en", "family", "family_fr", "order_fr", "iucn",
                                  "iucn_label")} | {"has_song": bool(e["songs"]), "n_cells": len(e["range"]["cells"])}
               for e in catalog.values()]

    class Handler(BaseHTTPRequestHandler):
        protocol_version = "HTTP/1.1"

        def log_message(self, fmt, *args):
            path = str(args[0]) if args else ""
            if not any(p in path for p in ("/audio/stream", "/api/m/image", "/photo/", "/gallery/")):
                super().log_message(fmt, *args)

        def send_json(self, obj, code=200):
            body = json.dumps(obj, ensure_ascii=False).encode()
            self.send_response(code)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def send_bytes(self, data, ctype, cache=True):
            """Réponse binaire avec prise en charge des requêtes Range (lecture audio sur iOS)."""
            total, rng = len(data), self.headers.get("Range")
            m = re.fullmatch(r"bytes=(\d*)-(\d*)", rng or "")
            if m and (m.group(1) or m.group(2)):
                a = int(m.group(1)) if m.group(1) else total - int(m.group(2))
                b = int(m.group(2)) if m.group(1) and m.group(2) else total - 1
                a, b = max(0, a), min(b, total - 1)
                if a > b:
                    self.send_response(416)
                    self.send_header("Content-Range", f"bytes */{total}")
                    self.send_header("Content-Length", "0")
                    return self.end_headers()
                self.send_response(206)
                self.send_header("Content-Range", f"bytes {a}-{b}/{total}")
                data = data[a:b + 1]
            else:
                self.send_response(200)
            self.send_header("Content-Type", ctype)
            self.send_header("Accept-Ranges", "bytes")
            self.send_header("Content-Length", str(len(data)))
            if cache:
                self.send_header("Cache-Control", "public, max-age=86400")
            self.end_headers()
            self.wfile.write(data)

        def read_json(self, limit):
            n = int(self.headers.get("Content-Length", 0))
            if n <= 0 or n > limit:
                raise ValueError("requête vide ou trop volumineuse")
            return json.loads(self.rfile.read(n))

        # ---------------------------------------------------------- GET
        def do_GET(self):
            url = urlparse(self.path)
            q = parse_qs(url.query)
            parts = [unquote(p) for p in url.path.strip("/").split("/")]
            try:
                if url.path == "/api/m/health":
                    return self.send_json({"ok": True, "species": len(catalog), "threshold": THRESHOLD})
                if url.path == "/api/m/species":
                    return self.send_json(summary)
                if parts[:3] == ["api", "m", "species"] and len(parts) == 4 and parts[3] in catalog:
                    return self.send_json(catalog[parts[3]])
                if parts[:3] == ["api", "m", "photo"] and len(parts) == 4 and parts[3] in catalog:
                    return self.send_bytes(thumb(catalog[parts[3]]["photo_key"], q.get("w", [0])[0]), "image/jpeg")
                if parts[:3] == ["api", "m", "gallery"] and len(parts) == 5 and parts[3] in catalog:
                    gal = catalog[parts[3]]["gallery"]
                    i = int(parts[4])
                    if 0 <= i < len(gal):
                        return self.send_bytes(thumb(gal[i]["bronze_key"], q.get("w", [0])[0]), "image/jpeg")
                if parts[:3] == ["api", "m", "song"] and len(parts) == 5 and parts[3] in catalog:
                    song = next((s for s in catalog[parts[3]]["songs"] if s["kind"] == parts[4]), None)
                    if song:
                        key = song["key"]
                        return self.send_bytes(bronze(key), "audio/wav" if key.endswith(".wav") else "audio/mpeg")
                if url.path == "/api/m/tests":
                    return self.send_json(list_tests())
                if url.path == "/api/m/dashboard":
                    return self.send_json(dashboard())
                if url.path == "/api/m/settings":
                    return self.send_json({"threshold": THRESHOLD, "min": THRESHOLD_RANGE[0], "max": THRESHOLD_RANGE[1]})
                if parts[:3] == ["api", "m", "tests"] and len(parts) >= 4 and re.fullmatch(r"[\w-]+", parts[3]):
                    meta = find_test(parts[3])
                    if not meta:
                        return self.send_error(404)
                    if len(parts) == 5 and parts[4] == "media":
                        return self.send_bytes(bronze(meta["media_key"]), meta["media_type"], cache=False)
                    return self.send_json(meta)
            except Exception as e:
                return self.send_json({"error": str(e)}, 500)
            self.send_error(404)

        # ---------------------------------------------------------- POST
        def do_POST(self):
            url = urlparse(self.path)
            try:
                if url.path == "/api/m/audio/stream":
                    return self.send_json(audio_stream(self.read_json(4 * 1024 * 1024)))
                if url.path == "/api/m/audio/file":
                    return self.send_json(audio_file(self.read_json(80 * 1024 * 1024)))
                if url.path == "/api/m/image":
                    return self.send_json(image(self.read_json(40 * 1024 * 1024)))
                if url.path == "/api/m/save":
                    return self.send_json(save(self.read_json(64 * 1024)))
                if url.path == "/api/m/settings":
                    value = set_threshold(self.read_json(4096).get("threshold"))
                    print(f"Seuil de confiance réglé à {value:.0%}")
                    return self.send_json({"threshold": value})
            except ValueError as e:
                return self.send_json({"error": str(e)}, 400)
            except Exception as e:
                return self.send_json({"error": f"{type(e).__name__}: {e}"}, 500)
            self.send_error(404)

    # ------------------------------------------------------------- traitements
    def audio_stream(req):
        s = sessions.get(req.get("session"), "audio", "Temps réel")
        sr = int(req.get("sr", 32000))
        if not 8000 <= sr <= 96000:
            raise ValueError("fréquence d'échantillonnage invalide")
        pcm = np.frombuffer(base64.b64decode(req.get("pcm", "")), dtype="<i2")
        if s["sr"] not in (None, sr):
            s["audio"], s["samples"] = [], 0
        s["sr"] = sr
        if s["samples"] < MAX_SESSION_S * sr:
            s["audio"].append(pcm.copy())
            s["samples"] += len(pcm)
        tail = np.concatenate(s["audio"][-64:]) if s["audio"] else np.zeros(0, "<i2")
        win = tail[-int(5 * sr):].astype(np.float32) / 32768.0
        t_end = s["samples"] / sr
        level = float(np.sqrt(np.mean(win[-int(sr / 2):] ** 2))) if len(win) else 0.0
        frame = int(sr / 4)
        peak = float(librosa.feature.rms(y=win, frame_length=frame, hop_length=frame // 2,
                                         center=False).max()) if len(win) >= frame else 0.0
        if len(win) < MIN_AUDIO_S * sr:
            status, top = "warmup", []
        elif peak < MIN_LEVEL:
            status, top = "quiet", []
        else:
            status, top = "ok", audio_model.predict(win, sr, k=3)
        new = []
        sp = top[0]["species"] if top and top[0]["prob"] >= THRESHOLD else None
        prev, n = s["streak"]
        s["streak"] = (sp, n + 1 if sp is not None and sp == prev else 1 if sp else 0)
        if sp and s["streak"][1] >= CONFIRM_WINDOWS:
            p = top[0]["prob"]
            if merge_detection(s, sp, p, max(0.0, t_end - 5), t_end):
                new.append(sp)
                log.add("Audio", "Temps réel", req["session"], sp, p, max(0.0, t_end - 5), t_end, position(req))
        return {"top": top, "level": level, "peak": round(peak, 4), "status": status, "seconds": round(t_end, 1), "new": new,
                "detections": detections_list(s), "threshold": THRESHOLD}

    def audio_file(req):
        s = sessions.get(req.get("session"), "audio", "Import")
        data = base64.b64decode(req.get("data", ""))
        y, sr = decode_audio(data, req.get("name"))
        if len(y) < sr * 0.5:
            raise ValueError("audio trop court (moins de 0,5 s)")
        y = y[: MAX_SESSION_S * sr]
        s.update(audio=[(np.clip(y, -1, 1) * 32767).astype("<i2")], sr=sr, samples=len(y), detections={},
                 name=req.get("name"))
        probs, starts = audio_model.windows(y, sr)
        timeline = []
        for p, t0 in zip(probs, starts):
            i = int(p.argmax())
            timeline.append({"t": round(t0, 2), "species": audio_model.classes[i], "conf": float(p[i])})
            if p[i] >= THRESHOLD:
                merge_detection(s, audio_model.classes[i], float(p[i]), t0, t0 + 5)
        for d in detections_list(s):
            log.add("Audio", "Import", req["session"], d["species"], d["conf"], d["segments"][0][0],
                    d["segments"][0][1], position(req))
        return {"duration": round(len(y) / sr, 2), "detections": detections_list(s), "timeline": timeline,
                "threshold": THRESHOLD}

    def image(req):
        source = "Temps réel" if req.get("mode") == "live" else "Import"
        s = sessions.get(req.get("session"), "image", source)
        data = base64.b64decode(req.get("data", ""))
        im, boxes = image_model.detect(data)
        kept = [b for b in boxes if b["conf"] >= THRESHOLD]
        if source == "Import":
            s["detections"] = {}
        new = []
        for b in kept:
            if merge_detection(s, b["species"], b["conf"]):
                new.append(b["species"])
                log.add("Image", source, req["session"], b["species"], b["conf"], pos=position(req))
        if kept or source == "Import" or s["image"] is None:
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=88)
            s["image"], s["boxes"] = buf.getvalue(), kept
        return {"width": im.width, "height": im.height, "boxes": kept, "best_below": next(
            (b for b in boxes if b["conf"] < THRESHOLD), None), "new": new, "detections": detections_list(s),
                "threshold": THRESHOLD}

    def save(req):
        sid = req.get("session")
        s = sessions.data.get(sid)
        if not s:
            raise ValueError("session expirée ou inconnue")
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        tid = f"{stamp}_{sid[:8]}"
        if s["kind"] == "audio":
            if not s["audio"]:
                raise ValueError("aucun audio à enregistrer")
            buf = io.BytesIO()
            sf.write(buf, np.concatenate(s["audio"]), s["sr"], format="WAV", subtype="PCM_16")
            media, ext, mtype = buf.getvalue(), "wav", "audio/wav"
        else:
            if not s["image"]:
                raise ValueError("aucune image à enregistrer")
            media, ext, mtype = s["image"], "jpg", "image/jpeg"
        base = f"{TESTS_PREFIX}{s['kind']}/{tid}"
        meta = {"id": tid, "kind": s["kind"], "source": s["source"], "created": s["created"],
                "saved": datetime.now().isoformat(timespec="seconds"), "name": s.get("name"),
                "duration": round(s["samples"] / s["sr"], 2) if s["kind"] == "audio" else None,
                "boxes": s["boxes"] if s["kind"] == "image" else None, "threshold": THRESHOLD,
                "detections": detections_list(s), "media_key": f"{base}.{ext}", "media_type": mtype}
        s3.put_object(Bucket=BRONZE, Key=f"{base}.{ext}", Body=media, ContentType=mtype)
        s3.put_object(Bucket=BRONZE, Key=f"{base}.json", Body=json.dumps(meta, ensure_ascii=False).encode(),
                      ContentType="application/json")
        return meta

    def dashboard():
        """Statistiques des détections (journal CSV + compteur) pour l'onglet Dashboard de l'app."""
        from collections import Counter, defaultdict
        from datetime import date, timedelta
        rows = log.rows()
        table = log.counter.table()
        sci2sp = {e["sci"]: sp for sp, e in catalog.items()}
        today = date.today()
        days = [(today - timedelta(days=i)).isoformat() for i in range(13, -1, -1)]
        by_day = {d: {"date": d, "audio": 0, "image": 0} for d in days}
        hours, modes, sources, fams, conf = [0] * 24, Counter(), Counter(), Counter(), []
        points, sessions = [], set()
        for r in rows:
            mode = "audio" if r["mode"] == "Audio" else "image"
            if r["date"] in by_day:
                by_day[r["date"]][mode] += 1
            try:
                hours[int(r["heure"][:2])] += 1
                conf.append(float(r["confiance_pct"]))
            except ValueError:
                pass
            modes[mode] += 1
            sources[r["source"]] += 1
            fams[r["famille"] or "?"] += 1
            sessions.add(r["session"])
            if r.get("latitude") and r.get("longitude"):
                points.append({"lat": float(r["latitude"]), "lon": float(r["longitude"]),
                               "species": sci2sp.get(r["nom_scientifique"], ""), "mode": mode,
                               "conf": float(r["confiance_pct"]), "when": f"{r['date']} {r['heure']}"})
        lo = THRESHOLD * 100                       # 4 classes de confiance entre le seuil et 100 %
        step = (100 - lo) / 4
        bins = [0, 0, 0, 0]
        for c in conf:
            bins[min(3, max(0, int((c - lo) // step)))] += 1
        bin_labels = [f"{lo + i * step:g}" for i in range(4)]
        bin_labels[-1] += "+"
        top = sorted((t for t in table if t["count"]), key=lambda t: (-t["count"], t["fr"]))
        iucn = Counter(catalog[t["species"]]["iucn"] for t in top)
        recent = [{"species": sci2sp.get(r["nom_scientifique"], ""), "mode": "audio" if r["mode"] == "Audio" else "image",
                   "source": r["source"], "conf": float(r["confiance_pct"] or 0), "when": f"{r['date']} {r['heure']}"}
                  for r in rows[-12:]][::-1]
        return {
            "total": len(rows), "sessions": len(sessions), "species_detected": len(top), "species_total": len(catalog),
            "avg_conf": round(sum(conf) / len(conf), 1) if conf else None,
            "modes": dict(modes), "sources": dict(sources),
            "by_day": list(by_day.values()), "by_hour": hours, "conf_bins": bins, "conf_labels": bin_labels,
            "families": [{"family": f, "count": n} for f, n in fams.most_common(8)],
            "iucn": dict(iucn), "top": top[:15], "points": points[-300:], "recent": recent,
            "files": {"journal": "data/igdad/detections.xlsx", "compteur": "data/igdad/compteur_especes.xlsx"},
        }

    def list_tests():
        out = []
        for page in s3.get_paginator("list_objects_v2").paginate(Bucket=BRONZE, Prefix=TESTS_PREFIX):
            for o in page.get("Contents", []):
                if o["Key"].endswith(".json"):
                    out.append(json.loads(bronze(o["Key"])))
        return sorted(out, key=lambda m: m["saved"], reverse=True)

    def find_test(tid):
        for kind in ("audio", "image"):
            try:
                return json.loads(bronze(f"{TESTS_PREFIX}{kind}/{tid}.json"))
            except s3.exceptions.NoSuchKey:
                continue
        return None

    return Handler


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--model", default=None, help="modèle audio (défaut : dernier runs/audio/*/best.pt)")
    ap.add_argument("--detector", default=None, help="poids YOLO (défaut : dernier runs/detect/*/weights/best.pt)")
    ap.add_argument("--host", default="0.0.0.0")
    ap.add_argument("--port", type=int, default=8000)
    args = ap.parse_args()
    ckpt = Path(args.model) if args.model else sorted((ROOT / "runs" / "audio").glob("*/best.pt"))[-1]
    det = Path(args.detector) if args.detector else sorted((ROOT / "runs" / "detect").glob("*/weights/best.pt"))[-1]
    load_settings()
    s3 = s3_client()
    catalog = {e["species"]: e for e in json.loads(
        s3.get_object(Bucket=GOLD, Key="species_catalog/species.json")["Body"].read())}
    print(f"Catalogue : {len(catalog)} espèces | audio : {ckpt} | image : {det}")
    log = DetectionLog(catalog)
    log.counter = SpeciesCounter(catalog, log.rows())
    ctx = {"s3": s3, "catalog": catalog, "audio": AudioModel(ckpt), "image": ImageModel(det),
           "sessions": Sessions(), "log": log}
    srv = ThreadingHTTPServer((args.host, args.port), make_handler(ctx))
    print(f"API Igdad : http://{args.host}:{args.port}  | journal : {EXPORT_DIR} | seuil : {THRESHOLD:.0%}")
    srv.serve_forever()


if __name__ == "__main__":
    main()
